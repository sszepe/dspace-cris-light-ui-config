"""
Management command: import_cris_layout

Parses the DSpace CRIS layout configuration XLS file and populates the
cris_layout database tables.

Usage
-----
    python manage.py import_cris_layout path/to/cris-layout-configuration.xls
    python manage.py import_cris_layout path/to/file.xls --clear   # wipe before import
    python manage.py import_cris_layout path/to/file.xls --entity Person,Publication
"""
from __future__ import annotations

import sys
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

try:
    import xlrd
except ImportError:
    xlrd = None  # type: ignore

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

from cris_layout.models import (
    CrisLayoutTab,
    CrisLayoutTab2Box,
    CrisLayoutBox,
    CrisLayoutBox2Metadata,
    CrisLayoutBox2Metrics,
    CrisLayoutBox2Vocabulary,
    CrisLayoutMetadataGroup,
    CrisLayoutTabPolicy,
    CrisLayoutBoxPolicy,
)


# ── helpers ───────────────────────────────────────────────────────────────────

def _yn(val: str) -> bool:
    return str(val).strip().lower() in ("y", "yes", "true", "1")


def _str(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    # xlrd returns floats for numeric cells — strip trailing .0 for IDs
    if s.endswith(".0") and s[:-2].lstrip("-").isdigit():
        s = s[:-2]
    return s


def _float(val) -> float:
    try:
        return float(val)
    except (TypeError, ValueError):
        return 0.0


def _int(val) -> int:
    try:
        return int(float(val))
    except (TypeError, ValueError):
        return 0


class WorkbookReader:
    """Thin abstraction over xlrd (.xls) and openpyxl (.xlsx)."""

    def __init__(self, path: str):
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"File not found: {path}")
        ext = p.suffix.lower()

        if ext in (".xls",):
            if xlrd is None:
                raise ImportError("xlrd is required for .xls files: pip install xlrd")
            self._wb = xlrd.open_workbook(path)
            self._mode = "xlrd"
        elif ext in (".xlsx", ".xlsm"):
            if openpyxl is None:
                raise ImportError("openpyxl is required for .xlsx files: pip install openpyxl")
            self._wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self._mode = "openpyxl"
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

    @property
    def sheet_names(self):
        if self._mode == "xlrd":
            return self._wb.sheet_names()
        return self._wb.sheetnames

    def sheet_rows(self, name: str) -> list[list[str]]:
        """Return all rows of a sheet as lists of stripped strings."""
        if self._mode == "xlrd":
            if name not in self._wb.sheet_names():
                return []
            sh = self._wb.sheet_by_name(name)
            return [[_str(sh.cell_value(r, c)) for c in range(sh.ncols)]
                    for r in range(sh.nrows)]
        else:
            if name not in self._wb.sheetnames:
                return []
            sh = self._wb[name]
            return [[_str(cell.value) for cell in row] for row in sh.iter_rows()]


# ── importers (one per sheet) ─────────────────────────────────────────────────

def import_tab(rows, entity_filter, stats):
    created = updated = 0
    for r in rows[1:]:  # skip header
        if len(r) < 6:
            continue
        entity, shortname, label, priority, leading, security = (
            r[0], r[1], r[2], r[3], r[4], r[5]
        )
        if not entity or not shortname:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        obj, new = CrisLayoutTab.objects.update_or_create(
            entity=entity, shortname=shortname,
            defaults=dict(
                label=label,
                priority=_int(priority),
                leading=_yn(leading),
                security=security or "PUBLIC",
            ),
        )
        if new:
            created += 1
        else:
            updated += 1
    stats["tab"] = (created, updated)


def import_tab2box(rows, entity_filter, stats):
    created = updated = 0
    # Delete existing for entity_filter scope before re-importing (no natural unique key)
    qs = CrisLayoutTab2Box.objects.all()
    if entity_filter:
        qs = qs.filter(entity__in=entity_filter)
    qs.delete()

    for r in rows[1:]:
        if len(r) < 6:
            continue
        entity, tab, row, row_style, cell_style, boxes = (
            r[0], r[1], r[2], r[3], r[4], r[5]
        )
        if not entity or not tab or not boxes:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        CrisLayoutTab2Box.objects.create(
            entity=entity, tab=tab, row=_float(row),
            row_style=row_style, cell_style=cell_style, boxes=boxes,
        )
        created += 1
    stats["tab2box"] = (created, 0)


def import_box(rows, entity_filter, stats):
    created = updated = 0
    for r in rows[1:]:
        if len(r) < 9:
            continue
        entity, collapsed, box_type, shortname, label, container, minor, security, style = (
            r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8]
        )
        if not entity or not shortname:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        obj, new = CrisLayoutBox.objects.update_or_create(
            entity=entity, shortname=shortname,
            defaults=dict(
                label=label,
                box_type=box_type or "METADATA",
                collapsed=_yn(collapsed),
                container=_yn(container),
                minor=_yn(minor),
                security=security or "PUBLIC",
                style=style,
            ),
        )
        if new:
            created += 1
        else:
            updated += 1
    stats["box"] = (created, updated)


def import_box2metadata(rows, entity_filter, stats):
    created = 0
    qs = CrisLayoutBox2Metadata.objects.all()
    if entity_filter:
        qs = qs.filter(entity__in=entity_filter)
    qs.delete()

    for r in rows[1:]:
        if len(r) < 16:
            r = r + [""] * (16 - len(r))
        entity, box, row, cell, ftype, metadata, value, bundle, label, lah, rendering, vi, rs, cs, sl, sv = (
            r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7],
            r[8], r[9], r[10], r[11], r[12], r[13], r[14], r[15],
        )
        if not entity or not box:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        CrisLayoutBox2Metadata.objects.create(
            entity=entity, box=box,
            row=_float(row), cell=_float(cell),
            field_type=ftype or "METADATA",
            metadata=metadata, value=value, bundle=bundle,
            label=label,
            label_as_heading=_yn(lah),
            rendering=rendering,
            values_inline=_yn(vi),
            row_style=rs, cell_style=cs,
            style_label=sl, style_value=sv,
        )
        created += 1
    stats["box2metadata"] = (created, 0)


def import_box2metrics(rows, entity_filter, stats):
    created = updated = 0
    for r in rows[1:]:
        if len(r) < 3:
            continue
        entity, box, metric_type = r[0], r[1], r[2]
        if not entity or not box:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        obj, new = CrisLayoutBox2Metrics.objects.update_or_create(
            entity=entity, box=box,
            defaults=dict(metric_type=metric_type),
        )
        if new:
            created += 1
        else:
            updated += 1
    stats["box2metrics"] = (created, updated)


def import_box2vocabulary(rows, entity_filter, stats):
    created = updated = 0
    for r in rows[1:]:
        if len(r) < 4:
            continue
        entity, box, vocab, metadata = r[0], r[1], r[2], r[3]
        if not entity or not box:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        obj, new = CrisLayoutBox2Vocabulary.objects.update_or_create(
            entity=entity, box=box,
            defaults=dict(vocabulary=vocab, metadata=metadata),
        )
        if new:
            created += 1
        else:
            updated += 1
    stats["box2vocab"] = (created, updated)


def import_metadatagroups(rows, entity_filter, stats):
    created = 0
    qs = CrisLayoutMetadataGroup.objects.all()
    if entity_filter:
        qs = qs.filter(entity__in=entity_filter)
    qs.delete()

    for r in rows[1:]:
        if len(r) < 10:
            r = r + [""] * (10 - len(r))
        entity, parent, ftype, metadata, value, bundle, label, rendering, sl, sv = (
            r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9]
        )
        if not entity or not parent or not metadata:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        CrisLayoutMetadataGroup.objects.create(
            entity=entity, parent=parent,
            field_type=ftype or "METADATA",
            metadata=metadata, value=value, bundle=bundle,
            label=label, rendering=rendering,
            style_label=sl, style_value=sv,
        )
        created += 1
    stats["metadatagroups"] = (created, 0)


def import_tabpolicy(rows, entity_filter, stats):
    created = updated = 0
    for r in rows[1:]:
        if len(r) < 4:
            r = r + [""] * (4 - len(r))
        entity, shortname, metadata, group = r[0], r[1], r[2], r[3]
        if not entity or not shortname:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        obj, new = CrisLayoutTabPolicy.objects.get_or_create(
            entity=entity, shortname=shortname, metadata=metadata,
            defaults=dict(group=group),
        )
        if new:
            created += 1
        else:
            updated += 1
    stats["tabpolicy"] = (created, updated)


def import_boxpolicy(rows, entity_filter, stats):
    created = updated = 0
    for r in rows[1:]:
        if len(r) < 4:
            r = r + [""] * (4 - len(r))
        entity, shortname, metadata, group = r[0], r[1], r[2], r[3]
        if not entity or not shortname:
            continue
        if entity_filter and entity not in entity_filter:
            continue
        obj, new = CrisLayoutBoxPolicy.objects.get_or_create(
            entity=entity, shortname=shortname, metadata=metadata,
            defaults=dict(group=group),
        )
        if new:
            created += 1
        else:
            updated += 1
    stats["boxpolicy"] = (created, updated)


# ── command ───────────────────────────────────────────────────────────────────

class Command(BaseCommand):
    help = "Import DSpace CRIS layout configuration from an XLS/XLSX file"

    def add_arguments(self, parser):
        parser.add_argument("xls_file", help="Path to cris-layout-configuration.xls or .xlsx")
        parser.add_argument(
            "--clear", action="store_true",
            help="Delete ALL existing layout data before importing",
        )
        parser.add_argument(
            "--entity", default="",
            help="Comma-separated list of entity types to import (default: all)",
        )

    def handle(self, *args, **options):
        path = options["xls_file"]
        clear = options["clear"]
        entity_filter = (
            {e.strip() for e in options["entity"].split(",") if e.strip()}
            if options["entity"] else None
        )

        self.stdout.write(f"Reading: {path}")
        try:
            reader = WorkbookReader(path)
        except (FileNotFoundError, ImportError, ValueError) as e:
            raise CommandError(str(e))

        stats: dict[str, tuple[int, int]] = {}

        with transaction.atomic():
            if clear:
                self.stdout.write(self.style.WARNING("Clearing all CRIS layout data…"))
                for Model in [
                    CrisLayoutTab, CrisLayoutTab2Box, CrisLayoutBox,
                    CrisLayoutBox2Metadata, CrisLayoutBox2Metrics,
                    CrisLayoutBox2Vocabulary, CrisLayoutMetadataGroup,
                    CrisLayoutTabPolicy, CrisLayoutBoxPolicy,
                ]:
                    Model.objects.all().delete()

            sheets = {
                "tab":                       import_tab,
                "tab2box":                   import_tab2box,
                "box":                       import_box,
                "box2metadata":              import_box2metadata,
                "box2metrics":               import_box2metrics,
                "box2hierarchicalvocabulary": import_box2vocabulary,
                "metadatagroups":            import_metadatagroups,
                "tabpolicy":                 import_tabpolicy,
                "boxpolicy":                 import_boxpolicy,
            }

            available = set(reader.sheet_names)
            for sheet_name, importer in sheets.items():
                if sheet_name not in available:
                    self.stdout.write(self.style.WARNING(f"  Sheet '{sheet_name}' not found — skipping"))
                    continue
                rows = reader.sheet_rows(sheet_name)
                importer(rows, entity_filter, stats)

        # ── Summary ──
        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS("Import complete:"))
        labels = {
            "tab": "Tabs", "tab2box": "Tab→Box mappings", "box": "Boxes",
            "box2metadata": "Box→Metadata fields", "box2metrics": "Box→Metrics",
            "box2vocab": "Box→Vocabularies", "metadatagroups": "Metadata group fields",
            "tabpolicy": "Tab policies", "boxpolicy": "Box policies",
        }
        for key, (created, updated) in stats.items():
            label = labels.get(key, key)
            self.stdout.write(f"  {label}: {created} created, {updated} updated")
